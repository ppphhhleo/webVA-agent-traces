#!/usr/bin/env python3
"""Extract the trial-one behavior coding sheet into one trace-linked JSON file."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.parse import urljoin
from urllib.request import urlopen

from openpyxl import load_workbook


DEFAULT_SHEET = "Coded Agent Traces Trial 1"
DEFAULT_CATALOG_URL = (
    "https://agent-traces-collection.s3.us-west-2.amazonaws.com/data/catalog.json"
)
ROUND_TOKEN = re.compile(
    r"(?P<start_marker>[A-Za-z])?\s*(?P<start>\d+)"
    r"(?:\s*[-–—]\s*(?P<end_marker>[A-Za-z])?\s*(?P<end>\d+))?"
)

MODEL_ALIASES = {
    "opus 4.8": "Claude Opus 4.8",
    "claude opus 4.8": "Claude Opus 4.8",
    "sonnet 5": "Claude Sonnet 5",
    "claude sonnet 5": "Claude Sonnet 5",
    "gpt 5.4": "GPT-5.4",
    "gpt-5.4": "GPT-5.4",
    "gpt 5.5": "GPT-5.5",
    "gpt-5.5": "GPT-5.5",
}
APP_ALIASES = {
    "data voyager": "Data Voyager",
    "datavoyager": "Data Voyager",
    "vitessce": "Vitessce",
    "lineup": "LineUp",
    "lit": "LIT",
    "sanddance": "SandDance",
    "gapminder": "Gapminder",
    "embedding atlas": "Embedding Atlas",
    "usgs": "USGS",
}
TASK_TYPE_ALIASES = {
    "low-level": "Low-level",
    "low level": "Low-level",
    "compound": "Compound",
    "high-level": "High-level",
    "high level": "High-level",
}


def clean_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def integer(value: Any, field: str) -> int:
    if isinstance(value, bool):
        raise ValueError(f"{field} must be an integer, got {value!r}")
    try:
        numeric = float(value)
    except (TypeError, ValueError) as error:
        raise ValueError(f"{field} must be an integer, got {value!r}") from error
    if not numeric.is_integer():
        raise ValueError(f"{field} must be an integer, got {value!r}")
    return int(numeric)


def number_or_text(value: Any) -> int | float | str | None:
    if value is None or clean_text(value) == "":
        return None
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        return int(value) if float(value).is_integer() else float(value)
    text = clean_text(value)
    try:
        numeric = float(text)
    except ValueError:
        return text
    return int(numeric) if numeric.is_integer() else numeric


def normalize(value: Any, aliases: dict[str, str], field: str) -> str:
    text = clean_text(value)
    normalized = aliases.get(text.lower())
    if not normalized:
        raise ValueError(f"Unknown {field}: {text!r}")
    return normalized


def parse_round_reference(raw_value: Any) -> dict[str, Any]:
    raw = clean_text(raw_value)
    if not raw:
        raise ValueError("Missing round reference")
    spans: list[dict[str, int]] = []
    warnings: list[str] = []
    for match in ROUND_TOKEN.finditer(raw):
        start = int(match.group("start"))
        end = int(match.group("end") or start)
        markers = {
            marker.upper()
            for marker in (match.group("start_marker"), match.group("end_marker"))
            if marker
        }
        for marker in sorted(markers - {"R"}):
            warnings.append(f"nonstandard round marker {marker}")
        if start > end:
            warnings.append(f"reversed range R{start}-R{end}")
            start, end = end, start
        spans.append({"start": start, "end": end})
    if not spans:
        raise ValueError(f"Could not parse round reference: {raw!r}")
    rounds = sorted({value for span in spans for value in range(span["start"], span["end"] + 1)})
    return {
        "raw": raw,
        "rounds": rounds,
        "spans": spans,
        "parse_warnings": sorted(set(warnings)),
    }


def fetch_json(url: str) -> dict[str, Any]:
    with urlopen(url, timeout=60) as response:
        return json.load(response)


def trace_round_count(catalog_url: str, trace: dict[str, Any]) -> int:
    trajectory = fetch_json(urljoin(catalog_url, trace["trajectory_url"]))
    return sum(
        event.get("event_type") == "model_message"
        for event in trajectory.get("events", [])
    )


def workbook_rows(path: Path, sheet_name: str) -> tuple[list[str], list[tuple[int, dict[str, Any]]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Missing sheet {sheet_name!r}; found {workbook.sheetnames}")
    sheet = workbook[sheet_name]
    iterator = sheet.iter_rows(values_only=True)
    headers = [clean_text(value) for value in next(iterator)]
    if len(headers) != len(set(headers)):
        raise ValueError("Sheet contains duplicate column headers")
    rows = [
        (row_number, dict(zip(headers, values)))
        for row_number, values in enumerate(iterator, start=2)
        if any(value is not None and clean_text(value) != "" for value in values)
    ]
    workbook.close()
    return headers, rows


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--sheet", default=DEFAULT_SHEET)
    parser.add_argument("--catalog-url", default=DEFAULT_CATALOG_URL)
    parser.add_argument(
        "--output",
        type=Path,
        default=Path(__file__).resolve().parents[1]
        / "data"
        / "annotations"
        / "agent_behaviors_trial1.json",
    )
    parser.add_argument("--workers", type=int, default=12)
    parser.add_argument("--expected-traces", type=int, default=120)
    args = parser.parse_args()

    required_headers = {
        "trace number id", "Trial", "Model", "App", "task id", "Task Type",
        "score (out of 1)", "trace id", "Task", "Theme", "Code", "Rounds", "Detail",
    }
    headers, rows = workbook_rows(args.workbook, args.sheet)
    missing_headers = required_headers - set(headers)
    if missing_headers:
        raise ValueError(f"Missing required columns: {sorted(missing_headers)}")

    catalog_document = fetch_json(args.catalog_url)
    catalog = catalog_document.get("traces", [])
    catalog_by_id = {clean_text(trace.get("trace_id")): trace for trace in catalog}
    catalog_base_url = args.catalog_url.rsplit("/", 1)[0] + "/"

    source_trace_rows: dict[str, tuple[int, dict[str, Any]]] = {}
    coding_rows: list[tuple[int, dict[str, Any]]] = []
    for row_number, row in rows:
        trace_id = clean_text(row.get("trace id"))
        if not trace_id:
            raise ValueError(f"Row {row_number} is missing trace id")
        if row.get("trace number id") is not None and clean_text(row.get("trace number id")):
            if trace_id in source_trace_rows:
                raise ValueError(f"Duplicate trace header for {trace_id} at row {row_number}")
            source_trace_rows[trace_id] = (row_number, row)
        if clean_text(row.get("Code")):
            coding_rows.append((row_number, row))

    if len(source_trace_rows) != args.expected_traces:
        raise ValueError(
            f"Expected {args.expected_traces} trace headers, found {len(source_trace_rows)}"
        )
    missing_catalog_ids = sorted(set(source_trace_rows) - set(catalog_by_id))
    if missing_catalog_ids:
        raise ValueError(f"Trace IDs absent from public catalog: {missing_catalog_ids}")

    normalized_headers: dict[str, dict[str, Any]] = {}
    labels_by_task: dict[int, list[str]] = defaultdict(list)
    for trace_id, (row_number, row) in source_trace_rows.items():
        trial = integer(row.get("Trial"), f"row {row_number} Trial")
        if trial != 1:
            raise ValueError(f"Expected trial 1 at row {row_number}, found {trial}")
        task_id = integer(row.get("task id"), f"row {row_number} task id")
        task_type = normalize(row.get("Task Type"), TASK_TYPE_ALIASES, "task type")
        labels_by_task[task_id].append(task_type)
        normalized_headers[trace_id] = {
            "source_row": row_number,
            "trace_number": integer(row.get("trace number id"), f"row {row_number} trace number"),
            "trial": trial,
            "source_model": normalize(row.get("Model"), MODEL_ALIASES, "model"),
            "app": normalize(row.get("App"), APP_ALIASES, "app"),
            "task_id": task_id,
            "source_task_type": task_type,
            "score": number_or_text(row.get("score (out of 1)")),
            "source_task_prompt": clean_text(row.get("Task")),
        }

    canonical_task_types = {
        task_id: Counter(labels).most_common(1)[0][0]
        for task_id, labels in labels_by_task.items()
    }
    selected_catalog = [catalog_by_id[trace_id] for trace_id in source_trace_rows]
    with ThreadPoolExecutor(max_workers=args.workers) as pool:
        round_counts = dict(zip(
            source_trace_rows,
            pool.map(lambda trace: trace_round_count(catalog_base_url, trace), selected_catalog),
        ))

    trace_index: list[dict[str, Any]] = []
    trace_metadata: dict[str, dict[str, Any]] = {}
    metadata_warnings: list[dict[str, Any]] = []
    for trace_id, source in normalized_headers.items():
        catalog_trace = catalog_by_id[trace_id]
        catalog_task_id = integer(catalog_trace.get("task_id"), f"catalog task id for {trace_id}")
        catalog_model = clean_text(catalog_trace.get("model"))
        if source["task_id"] != catalog_task_id:
            raise ValueError(
                f"Task mismatch for {trace_id}: workbook {source['task_id']} vs catalog {catalog_task_id}"
            )
        if source["source_model"] != catalog_model:
            raise ValueError(
                f"Model mismatch for {trace_id}: workbook {source['source_model']} vs catalog {catalog_model}"
            )
        canonical_type = canonical_task_types[source["task_id"]]
        metadata = {
            "trace_number": source["trace_number"],
            "catalog_trace_number": catalog_trace.get("trace_number"),
            "trace_id": trace_id,
            "trial": source["trial"],
            "model": catalog_model,
            "app": source["app"],
            "task_id": source["task_id"],
            "task_type": canonical_type,
            "score": source["score"],
            "task_prompt": clean_text(catalog_trace.get("task_prompt")) or source["source_task_prompt"],
            "round_count": round_counts[trace_id],
        }
        if source["source_task_type"] != canonical_type:
            metadata["source_task_type"] = source["source_task_type"]
            metadata_warnings.append({
                "trace_id": trace_id,
                "source_row": source["source_row"],
                "field": "task_type",
                "source_value": source["source_task_type"],
                "normalized_value": canonical_type,
            })
        trace_metadata[trace_id] = metadata
        trace_index.append(metadata.copy())
    trace_index.sort(key=lambda trace: trace["trace_number"])

    annotations: list[dict[str, Any]] = []
    seen_rows: set[tuple[Any, ...]] = set()
    for annotation_number, (row_number, row) in enumerate(coding_rows, start=1):
        trace_id = clean_text(row.get("trace id"))
        if trace_id not in trace_metadata:
            raise ValueError(f"Coding row {row_number} has no trace header: {trace_id}")
        theme = clean_text(row.get("Theme"))
        code = clean_text(row.get("Code"))
        detail = clean_text(row.get("Detail"))
        if not theme or not code or not detail:
            raise ValueError(f"Coding row {row_number} has a missing theme, code, or detail")
        round_reference = parse_round_reference(row.get("Rounds"))
        round_count = trace_metadata[trace_id]["round_count"]
        out_of_range = [value for value in round_reference["rounds"] if value > round_count]
        if out_of_range:
            round_reference["parse_warnings"].append(
                f"rounds exceed trace length {round_count}: {out_of_range}"
            )
        duplicate_key = (
            trace_id, theme, code, round_reference["raw"], detail,
        )
        if duplicate_key in seen_rows:
            raise ValueError(f"Duplicate coding row at source row {row_number}: {duplicate_key}")
        seen_rows.add(duplicate_key)
        metadata = trace_metadata[trace_id]
        annotations.append({
            "annotation_id": f"agent-trial1-{annotation_number:04d}",
            "source_row": row_number,
            "trace_number": metadata["trace_number"],
            "trace_id": trace_id,
            "model": metadata["model"],
            "app": metadata["app"],
            "task_id": metadata["task_id"],
            "task_type": metadata["task_type"],
            "theme": theme,
            "code": code,
            "round_reference": round_reference,
            "detail": detail,
        })

    annotations_by_trace = Counter(annotation["trace_id"] for annotation in annotations)
    traces_without_annotations = sorted(set(source_trace_rows) - set(annotations_by_trace))
    if traces_without_annotations:
        raise ValueError(f"Traces without annotations: {traces_without_annotations}")
    for trace in trace_index:
        trace["annotation_count"] = annotations_by_trace[trace["trace_id"]]

    theme_counts = Counter(annotation["theme"] for annotation in annotations)
    code_counts: dict[str, Counter[str]] = defaultdict(Counter)
    for annotation in annotations:
        code_counts[annotation["theme"]][annotation["code"]] += 1
    taxonomy = [
        {
            "theme": theme,
            "annotation_count": theme_counts[theme],
            "codes": [
                {"code": code, "annotation_count": count}
                for code, count in sorted(code_counts[theme].items())
            ],
        }
        for theme in sorted(theme_counts)
    ]
    parse_warnings = [
        {
            "annotation_id": annotation["annotation_id"],
            "trace_id": annotation["trace_id"],
            "source_row": annotation["source_row"],
            "warnings": annotation["round_reference"]["parse_warnings"],
        }
        for annotation in annotations
        if annotation["round_reference"]["parse_warnings"]
    ]

    source_sha256 = hashlib.sha256(args.workbook.read_bytes()).hexdigest()
    output = {
        "schema_version": "1.0",
        "dataset": "webva-agent-behaviors-trial1",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "source": {
            "workbook_filename": args.workbook.name,
            "workbook_sha256": source_sha256,
            "sheet": args.sheet,
        },
        "counts": {
            "traces": len(trace_index),
            "annotations": len(annotations),
            "themes": len(theme_counts),
            "codes": len({annotation["code"] for annotation in annotations}),
            "round_parse_warnings": len(parse_warnings),
            "metadata_normalizations": len(metadata_warnings),
        },
        "round_reference_method": {
            "raw": "The original spreadsheet value is preserved.",
            "rounds": "Expanded, sorted, unique one-based round numbers.",
            "spans": "Inclusive ranges in the order written by the coder.",
            "validation": "Parsed rounds are checked against the public trajectory's model-message count.",
        },
        "metadata_warnings": metadata_warnings,
        "round_parse_warnings": parse_warnings,
        "taxonomy": taxonomy,
        "trace_index": trace_index,
        "annotations": annotations,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(output, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps({
        "output": str(args.output),
        "traces": len(trace_index),
        "annotations": len(annotations),
        "themes": len(theme_counts),
        "codes": len({annotation["code"] for annotation in annotations}),
        "round_parse_warnings": len(parse_warnings),
        "metadata_normalizations": len(metadata_warnings),
        "workbook_sha256": source_sha256,
    }, indent=2))


if __name__ == "__main__":
    main()
